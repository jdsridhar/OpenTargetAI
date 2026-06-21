"""
Export Module for OpenTargetAI.
Exports prediction results to CSV, Excel, HTML, and PDF formats.
"""

import csv
import json
import logging
from datetime import datetime
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)


class ResultExporter:
    """
    Exports PredictionResult objects to various formats.
    """

    def __init__(self, prediction_result, query_smiles: str) -> None:
        self.result = prediction_result
        self.query_smiles = query_smiles
        self.timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    def export_csv(self, output_path: str) -> str:
        """Export predictions to CSV."""
        path = Path(output_path)
        with open(path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow([
                "Rank", "Target Name", "Gene Symbol", "UniProt", "Organism",
                "Target Class", "Confidence", "Confidence Score",
                "Supporting Ligands", "Mean Similarity", "Max Similarity",
                "Mean pChEMBL", "Activity Types"
            ])
            for i, pred in enumerate(self.result.predictions, 1):
                writer.writerow([
                    i, pred.target_name, pred.gene_symbol, pred.uniprot,
                    pred.organism, pred.target_class, pred.confidence_label,
                    f"{pred.confidence_score:.4f}", pred.n_supporting_ligands,
                    f"{pred.mean_similarity:.4f}", f"{pred.max_similarity:.4f}",
                    f"{pred.mean_pchembl:.2f}" if pred.mean_pchembl else "N/A",
                    "; ".join(pred.activity_types),
                ])
        logger.info(f"CSV exported to {path}")
        return str(path)

    def export_excel(self, output_path: str) -> str:
        """Export predictions to Excel (.xlsx)."""
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            raise RuntimeError("openpyxl is required for Excel export. Install it with: pip install openpyxl")

        wb = openpyxl.Workbook()

        # --- Summary Sheet ---
        ws_summary = wb.active
        ws_summary.title = "Summary"
        ws_summary.column_dimensions["A"].width = 25
        ws_summary.column_dimensions["B"].width = 50

        header_fill = PatternFill("solid", fgColor="1F4E79")
        header_font = Font(color="FFFFFF", bold=True)

        summary_data = [
            ("Query SMILES", self.query_smiles),
            ("Prediction Date", self.timestamp),
            ("Similar Compounds", self.result.n_similar_compounds),
            ("Targets Considered", self.result.n_targets_considered),
            ("Predictions", len(self.result.predictions)),
        ]
        for row_data in summary_data:
            ws_summary.append(row_data)

        # --- Predictions Sheet ---
        ws_preds = wb.create_sheet("Predictions")
        headers = [
            "Rank", "Target Name", "Gene Symbol", "UniProt", "Organism",
            "Target Class", "Confidence", "Confidence Score",
            "Supporting Ligands", "Mean Similarity", "Max Similarity",
            "Mean pChEMBL", "Activity Types",
        ]
        col_widths = [6, 40, 14, 12, 20, 16, 10, 16, 16, 14, 14, 12, 25]

        ws_preds.append(headers)
        for i, h in enumerate(headers, 1):
            cell = ws_preds.cell(row=1, column=i)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            ws_preds.column_dimensions[chr(64 + i)].width = col_widths[i - 1]

        conf_colors = {"High": "00B050", "Medium": "FFC000", "Low": "FF0000"}

        for rank, pred in enumerate(self.result.predictions, 1):
            row = [
                rank, pred.target_name, pred.gene_symbol, pred.uniprot,
                pred.organism, pred.target_class, pred.confidence_label,
                round(pred.confidence_score, 4), pred.n_supporting_ligands,
                round(pred.mean_similarity, 4), round(pred.max_similarity, 4),
                round(pred.mean_pchembl, 2) if pred.mean_pchembl else None,
                "; ".join(pred.activity_types),
            ]
            ws_preds.append(row)
            # Colour-code confidence
            conf_cell = ws_preds.cell(row=rank + 1, column=7)
            color = conf_colors.get(pred.confidence_label, "FFFFFF")
            conf_cell.fill = PatternFill("solid", fgColor=color)
            conf_cell.font = Font(bold=True)

        wb.save(output_path)
        logger.info(f"Excel exported to {output_path}")
        return output_path

    def export_html(self, output_path: str) -> str:
        """Export predictions to a standalone HTML report."""
        from jinja2 import Template

        template_str = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>OpenTargetAI Prediction Report</title>
<style>
  body { font-family: Arial, sans-serif; background: #0d1117; color: #c9d1d9; margin: 40px; }
  h1 { color: #58a6ff; } h2 { color: #58a6ff; border-bottom: 1px solid #30363d; padding-bottom: 4px; }
  table { border-collapse: collapse; width: 100%; margin-top: 16px; }
  th { background: #21262d; color: #58a6ff; padding: 10px; text-align: left; border: 1px solid #30363d; }
  td { padding: 8px 10px; border: 1px solid #21262d; }
  tr:nth-child(even) { background: #161b22; }
  tr:hover { background: #1c2128; }
  .high { color: #3fb950; font-weight: bold; }
  .medium { color: #d29922; font-weight: bold; }
  .low { color: #f85149; font-weight: bold; }
  .meta { background: #161b22; padding: 16px; border-radius: 8px; margin-bottom: 24px; }
  .smiles { font-family: monospace; background: #21262d; padding: 4px 8px; border-radius: 4px; }
</style>
</head>
<body>
<h1>🎯 OpenTargetAI — Prediction Report</h1>
<div class="meta">
  <p><strong>Query SMILES:</strong> <span class="smiles">{{ query_smiles }}</span></p>
  <p><strong>Generated:</strong> {{ timestamp }}</p>
  <p><strong>Similar Compounds:</strong> {{ n_similar }} &nbsp;|&nbsp;
     <strong>Targets Predicted:</strong> {{ n_preds }}</p>
</div>
<h2>Target Predictions</h2>
<table>
  <tr>
    <th>#</th><th>Target</th><th>Gene</th><th>UniProt</th>
    <th>Organism</th><th>Class</th><th>Confidence</th>
    <th>Score</th><th>Ligands</th><th>Sim̄</th><th>pChEMBL̄</th>
  </tr>
  {% for pred in predictions %}
  <tr>
    <td>{{ loop.index }}</td>
    <td>{{ pred.target_name }}</td>
    <td>{{ pred.gene_symbol }}</td>
    <td>{{ pred.uniprot }}</td>
    <td>{{ pred.organism }}</td>
    <td>{{ pred.target_class }}</td>
    <td class="{{ pred.confidence_label | lower }}">{{ pred.confidence_label }}</td>
    <td>{{ '%.4f' | format(pred.confidence_score) }}</td>
    <td>{{ pred.n_supporting_ligands }}</td>
    <td>{{ '%.3f' | format(pred.mean_similarity) }}</td>
    <td>{{ '%.2f' | format(pred.mean_pchembl) if pred.mean_pchembl else 'N/A' }}</td>
  </tr>
  {% endfor %}
</table>
<footer><p style="color:#484f58; margin-top:40px;">Generated by OpenTargetAI v1.0.0</p></footer>
</body>
</html>"""

        template = Template(template_str)
        html = template.render(
            query_smiles=self.query_smiles,
            timestamp=self.timestamp,
            n_similar=self.result.n_similar_compounds,
            n_preds=len(self.result.predictions),
            predictions=self.result.predictions,
        )
        Path(output_path).write_text(html, encoding="utf-8")
        logger.info(f"HTML report exported to {output_path}")
        return output_path

    def export_pdf(self, output_path: str) -> str:
        """Export predictions to PDF using reportlab."""
        try:
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.platypus import (
                SimpleDocTemplate, Table, TableStyle, Paragraph,
                Spacer, HRFlowable,
            )
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib import colors
            from reportlab.lib.units import cm
        except ImportError:
            raise RuntimeError("reportlab is required for PDF export. Install it with: pip install reportlab")

        doc = SimpleDocTemplate(
            output_path,
            pagesize=landscape(A4),
            leftMargin=1.5 * cm, rightMargin=1.5 * cm,
            topMargin=2 * cm, bottomMargin=2 * cm,
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "Title", parent=styles["Title"],
            textColor=colors.HexColor("#58a6ff"), fontSize=18,
        )
        normal = styles["Normal"]

        elements = [
            Paragraph("OpenTargetAI — Prediction Report", title_style),
            Spacer(1, 12),
            Paragraph(f"<b>Query SMILES:</b> {self.query_smiles}", normal),
            Paragraph(f"<b>Generated:</b> {self.timestamp}", normal),
            Paragraph(
                f"<b>Similar Compounds:</b> {self.result.n_similar_compounds}  "
                f"<b>Predictions:</b> {len(self.result.predictions)}", normal
            ),
            Spacer(1, 12),
            HRFlowable(width="100%", color=colors.HexColor("#30363d")),
            Spacer(1, 12),
        ]

        # Table data
        table_data = [[
            "#", "Target", "Gene", "UniProt", "Org.", "Class",
            "Conf.", "Score", "Lig.", "Sim̄", "pChEMBL̄",
        ]]
        for i, pred in enumerate(self.result.predictions, 1):
            table_data.append([
                str(i), pred.target_name[:28], pred.gene_symbol,
                pred.uniprot, pred.organism[:16], pred.target_class[:14],
                pred.confidence_label, f"{pred.confidence_score:.3f}",
                str(pred.n_supporting_ligands),
                f"{pred.mean_similarity:.3f}",
                f"{pred.mean_pchembl:.2f}" if pred.mean_pchembl else "N/A",
            ])

        col_widths = [1.0, 5.5, 2.2, 2.2, 2.8, 2.6, 1.8, 1.8, 1.3, 1.5, 2.0]
        col_widths = [w * cm for w in col_widths]

        table = Table(table_data, colWidths=col_widths, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#21262d")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#58a6ff")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [
                colors.HexColor("#161b22"), colors.HexColor("#0d1117")
            ]),
            ("TEXTCOLOR", (0, 1), (-1, -1), colors.HexColor("#c9d1d9")),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#30363d")),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))

        elements.append(table)
        doc.build(elements)
        logger.info(f"PDF exported to {output_path}")
        return output_path
